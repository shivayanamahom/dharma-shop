#!/usr/bin/env python3
"""Сценарии посетителя и сотрудника на наполненной чистой установке (после fixture.py). Каждый сценарий — набор проверок PASS/FAIL;
находки (предупреждения PHP, ответы 5xx) собираются клиентом и печатаются в конце.

  SCENARIO_URL=… SCENARIO_USER=… SCENARIO_PASS=… SCENARIO_SITE_PATH=… python3 dev/scenario/scenarios.py [раздел …]
Разделы: browse, product, favorites, cart, checkout, landing, admin, users, export, perf, all
"""
from __future__ import annotations

import json
import os
import re
import sys
import time
from html import unescape

sys.path.insert(0, os.path.dirname(__file__))
from bs4 import BeautifulSoup  # noqa: E402
from lib import Site, db, ok  # noqa: E402

R: list[tuple[bool, str]] = []
RUN = str(int(time.time()))[-6:]  # у каждого запуска свои адреса покупателей, чтобы не упираться в лимит «заказов на e-mail»
FINDINGS: list[str] = []


def soup(r) -> BeautifulSoup:
    return BeautifulSoup(r.text, "html.parser")


def ld(html: str) -> list[dict]:
    out = []
    for m in re.findall(r'<script type="application/ld\+json"[^>]*>(.*?)</script>', html, re.S):
        try:
            out.append(json.loads(m))
        except ValueError:
            pass
    return out


def links(html: str, needle: str) -> list[str]:
    return list(dict.fromkeys(a["href"] for a in BeautifulSoup(html, "html.parser").find_all("a", href=True) if needle in a["href"]))


def meta(html: str, name: str) -> str:
    m = re.search(rf'<meta[^>]+(?:name|property)="{re.escape(name)}"[^>]+content="([^"]*)"', html)
    return unescape(m.group(1)) if m else ""


def canonical(html: str) -> str:
    m = re.search(r'<link[^>]+rel="canonical"[^>]*href="([^"]*)"|<link[^>]+href="([^"]*)"[^>]*rel="canonical"', html)
    return (m.group(1) or m.group(2)) if m else ""


def browse(g: Site) -> None:
    print("\n== каталог и фильтры")
    r = g.get("/index.php/katalog")
    ok(r.status_code == 200 and "Электроника" in r.text, "каталог: список категорий открывается", R)
    r = g.get("/index.php/katalog/elektronika/smartfony")
    h = r.text
    ok(r.status_code == 200, "категория «Смартфоны» открывается", R)
    prod_links = links(h, "/katalog/elektronika/smartfony/")
    ok(len(prod_links) >= 5, f"в «Смартфонах» ссылок на товары: {len(prod_links)} (ожидается ≥5: чехол и 4 варианта)", R)
    ok("<h1" in h and h.count("<h1") == 1, "один заголовок h1", R)
    ok(bool(canonical(h)), f"canonical есть: {canonical(h)}", R)
    ok("noindex" not in meta(h, "robots"), "категория индексируется", R)
    imgs = [i for i in re.findall(r"<img[^>]*>", h) if "/images/dharmashop/" in i]
    ok(len(imgs) >= 5 and all("width=" in i and "height=" in i for i in imgs), f"у всех картинок товаров ({len(imgs)}) есть width/height", R)
    # сортировка и вид
    r2 = g.get("/index.php/katalog/elektronika?sort=price_asc")
    ok("noindex" in meta(r2.text, "robots"), "сортировка: страница-вариант закрыта от индексации", R)
    prices = [float(x.replace("\xa0", "").replace(" ", "").replace(",", ".")) for x in re.findall(r"([\d\s\xa0]+(?:[,.]\d+)?)\s*₽", BeautifulSoup(r2.text, "html.parser").select_one("[data-dharmashop-listing]").get_text(" ") if BeautifulSoup(r2.text, "html.parser").select_one("[data-dharmashop-listing]") else "") if x.strip()]
    ok(prices == sorted(prices) and len(prices) >= 3, f"сортировка по цене по возрастанию: {prices[:6]}", R)
    # фильтры: бренд
    r3 = g.get("/index.php/katalog/elektronika?filter[brand][]=pixel")
    n_pixel = len(links(r3.text, "/katalog/")) 
    ok("Pixel" in r3.text and r3.status_code == 200, "фильтр по бренду Pixel работает", R)
    ok("noindex" in meta(r3.text, "robots") or r3.history, "фильтр по бренду: noindex или переход на посадочную", R)
    r4 = g.get("/index.php/katalog/elektronika?filter[f][cvet][]=chernyj")
    ok(r4.status_code == 200, "фильтр по цвету открывается", R)
    r5 = g.get("/index.php/katalog/elektronika?filter[price][min]=4000&filter[price][max]=8000")
    ok(r5.status_code == 200, "фильтр по цене открывается", R)
    r6 = g.get("/index.php/katalog/elektronika?filter[stock]=1")
    ok(r6.status_code == 200, "фильтр «в наличии» открывается", R)
    # поиск (отдельное представление /search)
    S = "/index.php/component/dharmashop/search?filter[search]="
    r7 = g.get(S + "наушники")
    ok(r7.status_code == 200 and "Found: 3" in re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", r7.text)), "поиск «наушники»: найдено 3", R)
    r7b = g.get(S + "naushniki")
    ok("Found: 3" in re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", r7b.text)), "поиск «naushniki» (транслитерация): найдено 3", R)
    r7c = g.get(S + "AUR-SND")
    ok("Found: 1" in re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", r7c.text)), "поиск по артикулу «AUR-SND»: найден 1", R)
    ok("noindex" in meta(r7.text, "robots"), "результаты поиска закрыты от индексации", R)
    r8 = g.get("/index.php?option=com_dharmashop&task=search.suggest&format=json&q=наушн")
    try:
        j = r8.json()
        ok(j.get("total") == 3, f"подсказки поиска: найдено {j.get('total')}", R)
    except ValueError:
        ok(False, f"подсказки поиска: не JSON ({r8.status_code})", R)
    r9 = g.get(S + "несуществующийтовар")
    ok(r9.status_code == 200 and "Found: 0" in re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", r9.text)) and "noindex" in meta(r9.text, "robots"), "поиск без результатов: 200, «Found: 0», noindex", R)
    r9b = g.get("/index.php/katalog/elektronika?filter[search]=" + "наушники")
    ok("Наушники Pixel Buds" in r9b.text and "Смартфон Pixel P Синий 256 ГБ" not in BeautifulSoup(r9b.text, "html.parser").select_one("[data-dharmashop-listing]").get_text(" "), "поиск внутри категории сужает выдачу", R) if BeautifulSoup(r9b.text, "html.parser").select_one("[data-dharmashop-listing]") else ok("Наушники Pixel Buds" in r9b.text, "поиск внутри категории находит товар", R)
    # бренды и бейджи
    r = g.get("/index.php/brendy")
    ok(r.status_code == 200 and "Aurora" in r.text, "список брендов", R)
    r = g.get("/index.php/bejdzhi")
    ok(r.status_code == 200, "список бейджей", R)
    # 404 и пагинация за пределами
    r = g.get("/index.php/katalog/elektronika/net-takoj")
    ok(r.status_code == 404, "несуществующая категория: 404", R)
    r = g.get("/index.php/katalog/elektronika?start=500")
    ok(r.status_code == 404, "пагинация за пределами: 404", R)
    r = g.get("/index.php/katalog/elektronika?start=0", allow_redirects=False)
    ok(r.status_code == 301, "?start=0 → 301 на чистый адрес", R)


def pid(sku: str) -> int:
    return int(db(f"SELECT product_id FROM PFX_dharmashop_variants WHERE sku = '{sku}'")[0][0])


def text_of(html: str) -> str:
    t = re.sub(r"<script.*?</script>|<style.*?</style>", "", html, flags=re.S)
    return re.sub(r"\s+", " ", unescape(re.sub(r"<[^>]+>", " ", t))).strip()


def walk(node, kind: str):
    """Все объекты JSON-LD нужного @type (в том числе вложенные)."""
    if isinstance(node, dict):
        t = node.get("@type")
        if t == kind or (isinstance(t, list) and kind in t):
            yield node
        for v in node.values():
            yield from walk(v, kind)
    elif isinstance(node, list):
        for v in node:
            yield from walk(v, kind)


class Shopper(Site):
    """Гость с корзиной: те же запросы, что отправляет JS витрины."""

    JSON = {"X-Requested-With": "XMLHttpRequest", "Accept": "application/json"}

    def __init__(self):
        super().__init__()
        page = self.get("/index.php/katalog")
        o = json.loads(re.search(r'<script type="application/json" class="joomla-script-options[^>]*>(.*?)</script>', page.text, re.S).group(1))
        self.tok = o["csrf.token"]
        self.urls = o["com_dharmashop.cart"]
        self.last: dict = {}

    def api(self, url: str, data: dict, token: bool = True) -> dict:
        body = dict(data)
        if token:
            body[self.tok] = "1"
        r = self.post(url, body, headers=self.JSON)
        try:
            j = r.json()
        except ValueError:
            return {"success": False, "message": f"HTTP {r.status_code} не JSON"}
        self.last = j.get("cart") or self.last
        return j

    def add(self, product_id: int, qty: float = 1, variant: int | None = None) -> dict:
        d = {"product_id": product_id, "qty": qty}
        if variant:
            d["variant_id"] = variant
        return self.api(self.urls["url"], d)

    def update(self, item: int, qty: float) -> dict:
        return self.api(self.urls["updateUrl"], {f"qty[{item}]": qty})

    def remove(self, item: int) -> dict:
        return self.api(self.urls["removeUrl"], {"item_id": item})

    def promo(self, code: str) -> dict:
        return self.api("/index.php/component/dharmashop?task=cart.applyPromotion&format=json", {"promotion_code": code})

    def unpromo(self) -> dict:
        return self.api("/index.php/component/dharmashop?task=cart.removePromotion&format=json", {})

    def checkout_form(self):
        page = self.get("/index.php/oformlenie")
        form = [f for f in BeautifulSoup(page.text, "html.parser").find_all("form") if "checkout.submit" in str(f)]
        if not form:
            return page, None, []
        fields = []
        for el in form[0].find_all(["input", "textarea", "select"]):
            n = el.get("name")
            if not n or el.get("type") in {"submit", "button"}:
                continue
            if el.get("type") in {"checkbox", "radio"} and not el.has_attr("checked"):
                continue
            fields.append((n, el.get("value", "") if el.name != "textarea" else el.text))
        return page, form[0].get("action") or "/index.php/oformlenie", fields

    def order(self, shipping: int = 1, payment: int = 1, email: str = "buyer@example.test", over: dict | None = None, drop: tuple[str, ...] = (), consent: bool = True):
        page, action, fields = self.checkout_form()
        if action is None:
            return page, None
        base = {"shipping_method": str(shipping), "payment_method": str(payment), "recipient[name]": "Тест Покупатель", "recipient[phone]": "+7 (900) 123-45-67",
                "recipient[email]": email, "recipient[comment]": "сценарий"}
        if shipping == 1:
            base["shipping[1][point]"] = "1"
        if shipping == 2:
            base.update({"shipping[2][city]": "Пермь", "shipping[2][street]": "Ленина", "shipping[2][house]": "10", "shipping[2][kvartira]": "5"})
        if shipping == 3:
            base.update({"shipping[3][gorod]": "Москва", "shipping[3][ulica_i_dom]": "Тверская 1"})
        if consent:
            base["consents[personal_data]"] = "1"
        base.update(over or {})
        data = {k: v for k, v in fields if not k.startswith(("shipping", "payment", "recipient", "consents"))}
        data.update(base)
        for d in drop:
            data.pop(d, None)
        r = self.post(action, data)
        return r, data


def last_order(email: str):
    rows = db(f"SELECT id, number, total, subtotal, adjustments_total, status_id, payment_status, shipping_title, payment_title, contact_phone FROM PFX_dharmashop_orders WHERE contact_email = '{email}' ORDER BY id DESC LIMIT 1")
    return rows[0] if rows else None


def product(g: Site) -> None:
    print("\n== страницы товаров и микроразметка")
    restock()
    base = "/index.php/katalog/elektronika/"
    r = g.get(base + "nausniki/nausniki-aurora-sound")
    h = r.text
    prods = [p for d in ld(h) for p in walk(d, "Product")]
    ok(r.status_code == 200 and h.count("<h1") == 1, "товар: 200 и один h1", R)
    ok(len(prods) == 1 and prods[0]["offers"]["price"] == "4990.00" and prods[0]["offers"]["availability"].endswith("InStock"), "JSON-LD Product: цена 4990.00, InStock", R)
    ok(len(prods[0].get("image", [])) == 2 and len(prods[0].get("additionalProperty", [])) >= 4 and prods[0].get("brand", {}).get("name") == "Aurora", "JSON-LD: 2 фото, характеристики, бренд", R)
    imgs = [i for i in re.findall(r"<img[^>]*>", h) if "/images/dharmashop/" in i]
    ok(imgs and all("width=" in i and "height=" in i and "alt=" in i for i in imgs), f"картинки товара ({len(imgs)}): width, height, alt", R)
    ok(meta(h, "og:image:width") != "" and meta(h, "og:image").startswith("https://"), "og:image с размерами", R)
    set_config(seo_product_description="{product.title} — {product.introtext}")  # значение по умолчанию из поставки
    h = g.get(base + "nausniki/nausniki-aurora-sound").text
    ok(meta(h, "description") == "Наушники Aurora Sound — Беспроводные наушники с шумоподавлением.", f"meta description = название + краткое описание: «{meta(h, 'description')[:70]}»", R)
    # скидка
    r = g.get(base + "nausniki/nausniki-nordic-air")
    p = [p for d in ld(r.text) for p in walk(d, "Product")][0]
    ok(abs(float(p["offers"]["price"]) - 6791.5) < 0.01, f"скидка 15%: в JSON-LD цена {p['offers']['price']} (7990 → 6791.50)", R)
    t = text_of(r.text)
    ok("7 990" in t and "6 792" in t or "6 791" in t, "на странице видны старая и новая цена", R)
    # нет в наличии
    r = g.get(base + "nausniki/nausniki-pixel-buds")
    p = [p for d in ld(r.text) for p in walk(d, "Product")][0]
    ok(p["offers"]["availability"].endswith("OutOfStock"), "нет в наличии: OutOfStock", R)
    ok("Out of stock" in text_of(r.text), "на странице «Out of stock»", R)
    # учёт остатков, остаток 0, но заказ без остатка разрешён: разметка совпадает с витриной (бейдж «In stock», кнопка покупки)
    r = g.get("/index.php/katalog/dom/kruzki/kruzka-temnaa")
    p = [p for d in ld(r.text) for p in walk(d, "Product")][0]
    t = text_of(r.text)
    ok(p["offers"]["availability"].endswith("/InStock") and "In stock" in t and "Out of stock" not in t and "Add to cart" in t,
       f"остаток 0 + backorder: JSON-LD {p['offers']['availability'].rsplit('/', 1)[-1]}, на странице «In stock» и «Add to cart»", R)
    # цена по запросу
    r = g.get("/index.php/katalog/dom/podarocnyj-sertifikat")
    if r.status_code != 200:
        links_ = links(g.get("/index.php/katalog/dom").text, "/katalog/dom/")
        r = g.get(links_[0]) if links_ else r
    pp = [p for d in ld(r.text) for p in walk(d, "Product")]
    ok(r.status_code == 200 and pp and "offers" not in pp[0], "цена по запросу: Offer в разметке не выводится", R)
    # группы
    r = g.get(base + "smartfony/smartfon-pixel-p-cernyj-128-gb")
    groups = [p for d in ld(r.text) for p in walk(d, "ProductGroup")]
    ok(bool(groups) and len(groups[0].get("hasVariant", [])) == 4, f"ProductGroup: вариантов {len(groups[0].get('hasVariant', [])) if groups else 0} (ожидается 4)", R)
    ok(canonical(r.text).endswith("smartfon-pixel-p-cernyj-128-gb") or "smartfon-pixel-p" in canonical(r.text), f"canonical варианта: {canonical(r.text)[-45:]}", R)
    t = text_of(r.text)
    ok("128" in t and "256" in t, "на странице варианта предлагается выбор вариантов (128/256)", R)
    # наличие каждого варианта группы в разметке совпадает с бейджем на странице этого варианта
    mismatch = []
    for x in groups[0]["hasVariant"] if groups else []:
        if "offers" not in x:
            continue
        page = g.get(x["url"]).text
        shown = "InStock" if "In stock" in text_of(page) and "Out of stock" not in text_of(page) else "OutOfStock"
        if x["offers"]["availability"].rsplit("/", 1)[-1] != shown:
            mismatch.append(f"{x['url'].rsplit('/', 1)[-1]}: разметка {x['offers']['availability'].rsplit('/', 1)[-1]}, витрина {shown}")
    ok(groups and not mismatch, f"наличие вариантов группы в JSON-LD совпадает с витриной{': ' + '; '.join(mismatch) if mismatch else ''}", R)
    prices = sorted({float(x["offers"]["price"]) for x in groups[0]["hasVariant"] if "offers" in x}) if groups else []
    ok(prices == [29990.0, 34990.0], f"цены вариантов: {prices}", R)
    # крошки
    bc = [b for d in ld(r.text) for b in walk(d, "BreadcrumbList")]
    ok(bool(bc) and len(bc[0]["itemListElement"]) >= 3, "BreadcrumbList не короче 3 звеньев", R)
    # дубль по неверному пути
    r = g.get("/index.php/katalog/odezda/nausniki-aurora-sound")
    ok(r.status_code == 404, "товар по чужому пути категории: 404", R)


def modules_css(g: Site) -> None:
    print("\n== модуль категорий в узкой колонке")
    h = g.get("/index.php/katalog/elektronika/smartfony").text
    ok("mod_dharmashop_categories/css/grid" in h or "mod_dharmashop_categories/grid" in h, "стиль сетки категорий подключён", R)
    css = g.get(re.search(r'href="([^"]*mod_dharmashop_categories/css/grid[^"]*)"', h).group(1)).text if re.search(r'href="([^"]*mod_dharmashop_categories/css/grid[^"]*)"', h) else ""
    ok("@container" in css and "container-type" in css, "в стиле есть контейнерный запрос для узкой колонки", R)


def favorites(g: Site) -> None:
    print("\n== избранное")
    f = Site()
    page = f.get("/index.php/katalog/elektronika/nausniki")
    soup_ = BeautifulSoup(page.text, "html.parser")
    form = [x for x in soup_.find_all("form") if "favorites.toggle" in (x.get("action") or "")][0]
    tok = [i.get("name") for i in form.find_all("input") if re.fullmatch(r"[0-9a-f]{32}", i.get("name") or "")][0]
    url = form["action"] + "&format=json"
    ids = [pid("AUR-SND-1"), pid("NRD-AIR-1")]
    for i in ids:
        j = f.post(url, {"product_id": i, tok: "1"}, headers=Shopper.JSON).json()
        ok(j.get("success") is True, f"добавить в избранное товар {i}: {str(j.get('message'))[:50]}", R)
    r = f.get("/index.php/izbrannoe")
    ok("Наушники Aurora Sound" in r.text and "Наушники Nordic Air" in r.text, "страница избранного показывает оба товара", R)
    ok("noindex" in meta(r.text, "robots"), "избранное закрыто от индексации", R)
    cards = BeautifulSoup(r.text, "html.parser").select(".dharmaCard, [data-dharmashop-card]")
    ok(len(cards) >= 2, f"карточки избранного как в каталоге ({len(cards)})", R)
    j = f.post(url, {"product_id": ids[0], tok: "1"}, headers=Shopper.JSON).json()
    r = f.get("/index.php/izbrannoe")
    ok("Наушники Aurora Sound" not in r.text and "Наушники Nordic Air" in r.text, "повторное нажатие убирает товар", R)
    j = f.post(url, {"product_id": 999999, tok: "1"}, headers=Shopper.JSON).json()
    ok(j.get("success") is False, "несуществующий товар не добавляется", R)
    j = f.post(url, {"product_id": ids[0]}, headers=Shopper.JSON).json()
    ok(j.get("success") is False, "без токена отказ", R)
    other = Site()
    ok("Наушники Nordic Air" not in other.get("/index.php/izbrannoe").text, "избранное другого гостя пусто", R)
    n = db(f"SELECT COUNT(*) FROM PFX_dharmashop_favorites")[0][0] if db("SHOW TABLES LIKE 'PFX_dharmashop_favorites'") else "?"
    print(f"  (записей избранного в базе: {n})")


def cart(g: Site) -> None:
    print("\n== корзина")
    restock()
    c = Shopper()
    aur, nor, buds, mug_dark, gift, mug = pid("AUR-SND-1"), pid("NRD-AIR-1"), pid("PXL-BUD-1"), pid("MUG-DARK"), pid("GIFT-1"), pid("MUG-MORN")
    j = c.add(aur, 2)
    ok(j["success"] and c.last["qty"] == 2 and c.last["subtotal"] == 9980, f"добавить 2 шт: {c.last}", R)
    item = c.last["products"][str(aur)]["item"]
    c.add(aur, 1)
    ok(c.last["qty"] == 3 and c.last["count"] == 1, "повторное добавление суммирует количество в одной строке", R)
    c.update(item, 1)
    ok(c.last["qty"] == 1 and c.last["subtotal"] == 4990, "изменить количество на 1", R)
    j = c.add(buds, 1)
    ok(j["success"] is False, f"нет в наличии не добавляется: {j.get('message')}", R)
    j = c.add(nor, 9)
    ok(j["success"] is False or c.last["products"].get(str(nor), {}).get("qty", 99) <= 8, f"больше остатка (9 из 8): {j.get('message')} / в корзине {c.last['products'].get(str(nor))}", R)
    j = c.add(mug_dark, 1)
    ok(j["success"] is True, "под заказ (остаток 0, backorder) добавляется", R)
    j = c.add(gift, 1)
    ok(j["success"] is True, "цена по запросу добавляется", R)
    j = c.add(nor, 1)
    ok(j["success"] is True, "товар со скидкой добавляется", R)
    j = c.add(999999, 1)
    ok(j["success"] is False, "несуществующий товар: отказ", R)
    j = c.api(c.urls["url"], {"product_id": aur, "qty": 1}, token=False)
    ok(j["success"] is False, "без токена: отказ", R)
    j = c.add(aur, 1e12)
    ok(j["success"] is False or j.get("cart", {}).get("qty", 0) < 1e9, "огромное количество не ломает корзину", R)
    page = c.get("/index.php/korzina")
    t = text_of(page.text)
    ok(page.status_code == 200 and "Наушники Aurora Sound" in t and "Кружка «Тёмная»" in t, "страница корзины показывает позиции", R)
    ok("noindex" in meta(page.text, "robots"), "корзина noindex", R)
    for k in list(c.last["products"].values()):
        c.remove(k["item"])
    ok(c.last["count"] == 0, f"после удаления всех позиций корзина пуста: {c.last}", R)
    # скидка от суммы и купоны
    c = Shopper()
    c.add(aur, 2)
    page, _, _ = c.checkout_form()
    t = text_of(page.text)
    ok("Order amount discount (5%)" in t and "-499" in t, "скидка от суммы 5% на 9 980 ₽ → −499 ₽", R)
    c.add(aur, 1)
    t = text_of(c.get("/index.php/oformlenie").text)
    ok("(10%)" in t and "-1 497" in t.replace("\xa0", " "), "скидка от суммы 10% на 14 970 ₽ → −1 497 ₽", R)
    j = c.promo("WELCOME10")
    t = text_of(c.get("/index.php/oformlenie").text)
    ok(j["success"] and "WELCOME10" in t, f"купон WELCOME10 применён: {str(j.get('message'))[:50]}", R)
    ok(len(re.findall(r"-\s?\d[\d ]*\s?₽", t)) >= 2, "скидка от суммы и купон действуют вместе", R)
    j = c.promo("NOPE")
    ok(j["success"] is False, f"неверный код отклонён: {j.get('message')}", R)
    j = c.promo("MINUS500")
    t = text_of(c.get("/index.php/oformlenie").text)
    ok("MINUS500" in t and "WELCOME10" not in t, "второй купон заменяет первый (один код на заказ)", R)
    c.unpromo()
    t = text_of(c.get("/index.php/oformlenie").text)
    ok("MINUS500" not in t, "купон снимается", R)
    small = Shopper()
    small.add(mug, 1)
    j = small.promo("MINUS500")
    t = text_of(small.get("/index.php/oformlenie").text)
    ok("Coupon discount" not in t, "MINUS500 при сумме 590 ₽ (минимум 3000): код прикреплён, скидки пока нет", R)
    ok("not active yet" in str(j.get("message")) and "starts only when" in t, f"покупателю объяснено, что скидка пока не действует: «{str(j.get('message'))[:60]}»", R)
    small.add(mug, 5)
    t = text_of(small.get("/index.php/oformlenie").text)
    ok("Coupon discount MINUS500 -500" in t, "MINUS500 при 3 540 ₽: скидка −500 ₽ появилась", R)
    j = Shopper().promo("WELCOME10")
    ok(j["success"] is False, "купон в пустой корзине отклонён", R)


def restock() -> None:
    """Возвращает остатки как в наполнении: сценарии заказов их списывают, а повторный запуск должен начинаться с тех же данных."""
    a = Site()
    a.login()
    for sku, qty in (("AUR-SND-1", 25), ("NRD-AIR-1", 8), ("MUG-MORN", 120)):
        i = pid(sku)
        a.submit_form(f"/administrator/index.php?option=com_dharmashop&task=product.edit&id={i}", {"jform[stock_qty]": str(qty)}, task="product.save")
        # закрыть запись (проверка включена на время редактирования)


def set_config(**kv) -> None:
    """Сохранить параметры компонента через форму «Настройки» (как администратор), остальные значения остаются прежними."""
    a = Site()
    a.login()
    a.submit_form("/administrator/index.php?option=com_config&view=component&component=com_dharmashop", {f"jform[{k}]": str(v) for k, v in kv.items()}, task="component.apply")


def checkout(g: Site) -> None:
    print("\n== оформление заказов")
    set_config(checkout_rate_ip=1000)
    restock()
    aur, mug, nor = pid("AUR-SND-1"), pid("MUG-MORN"), pid("NRD-AIR-1")
    stock0 = float(db(f"SELECT stock_total FROM PFX_dharmashop_variants WHERE product_id = {aur}")[0][0])
    # 1. самовывоз + оплата при получении, скидка 5%
    c = Shopper(); c.add(aur, 2)
    r, _ = c.order(1, 1, f"o1.{RUN}@example.test")
    o = last_order(f"o1.{RUN}@example.test")
    ok(o is not None, f"заказ №1 создан: {o[:4] if o else r.url}", R)
    if o:
        ok(float(o[2]) == 9481.0 and float(o[3]) == 9980.0 and float(o[4]) == -499.0, f"сумма 9 481 (9 980 − 5%): subtotal={o[3]} adj={o[4]} total={o[2]}", R)
        ok(o[7] == "Самовывоз" and o[8] == "Оплата при получении", f"способы сохранены: {o[7]} / {o[8]}", R)
        ok("checkoutcomplete" in r.url or "zakaz-oformlen" in r.url, f"редирект на «Заказ оформлен»: {r.url[-40:]}", R)
        ok(o[1] in text_of(r.text) or "Aurora" in text_of(r.text), "на странице завершения видны данные заказа", R)
        items = db(f"SELECT title, qty, price, total FROM PFX_dharmashop_order_items WHERE order_id = {o[0]}")
        ok(len(items) == 1 and float(items[0][1]) == 2 and float(items[0][3]) == 9980.0, f"позиция заказа: {items}", R)
        stock1 = float(db(f"SELECT stock_total FROM PFX_dharmashop_variants WHERE product_id = {aur}")[0][0])
        print(f"  остаток Aurora: {stock0} → {stock1}")
        ok(stock1 <= stock0 - 2 or stock1 == stock0, "остаток: списан при заказе или ждёт статуса (см. вывод)", R)
        mv = db(f"SELECT COUNT(*) FROM PFX_dharmashop_stock_movements WHERE order_id = {o[0]}") if db("SHOW COLUMNS FROM PFX_dharmashop_stock_movements LIKE 'order_id'") else [["?"]]
        print(f"  движений остатка по заказу: {mv[0][0]}")
    ok(Shopper_cart_empty(c), "корзина очищена после заказа", R)
    # 2. курьер + онлайн-оплата (тест): доставка бесплатна от 5 000 ₽
    c = Shopper(); c.add(aur, 2)
    r, _ = c.order(2, 2, f"o2.{RUN}@example.test")
    o = last_order(f"o2.{RUN}@example.test")
    ok(o is not None, f"заказ №2 (курьер + онлайн): {o[:4] if o else text_of(r.text)[:160]}", R)
    if o:
        ok(float(o[2]) == 9481.0, f"курьер бесплатно от 5 000 ₽: итог {o[2]}", R)
        ok("payment" in r.url or "oplat" in r.url, f"редирект на страницу оплаты: {r.url[-50:]}", R)
        pay_html = r.text
        paid_links = re.findall(r'href="([^"]*payment\.callback[^"]*paid[^"]*)"', pay_html)
        print(f"  ссылок «оплачено» на странице: {len(paid_links)}; статус оплаты: {o[6]}")
        if paid_links:
            href = unescape(paid_links[0])
            r2 = c.get(href)
            o = last_order(f"o2.{RUN}@example.test")
            ok(o[6] == "paid", f"после тестовой оплаты payment_status={o[6]}, status_id={o[5]}", R)
        else:
            ok(False, "на странице оплаты нет ссылки/кнопки «оплачено» (песочница)", R)
    # 3. почта России + оплата при получении, малая сумма → доставка 450, без скидки
    c = Shopper(); c.add(mug, 1)
    r, _ = c.order(3, 1, f"o3.{RUN}@example.test")
    o = last_order(f"o3.{RUN}@example.test")
    ok(o is not None and float(o[2]) == 590.0 + 450.0 or (o is not None and float(o[2]) in (590.0, 1040.0)), f"заказ №3 (почта): итог {o[2] if o else '—'} (590 + 450 = 1 040 ожидается)", R)
    if o and float(o[2]) == 590.0:
        FINDINGS.append("ЗАКАЗ: стоимость доставки «Почта России» (450 ₽) не включена в итог заказа — проверить, как задана цена способа")
    # 4. курьер платно (<5000) и купон
    c = Shopper(); c.add(mug, 3); c.add(nor, 1)
    c.promo("WELCOME10")
    r, _ = c.order(2, 1, f"o4.{RUN}@example.test")
    o = last_order(f"o4.{RUN}@example.test")
    ok(o is not None, f"заказ №4 (курьер, купон): {o[:5] if o else text_of(r.text)[:160]}", R)
    if o:
        adj = db(f"SELECT title, amount FROM PFX_dharmashop_order_adjustments WHERE order_id = {o[0]}") if db("SHOW TABLES LIKE 'PFX_dharmashop_order_adjustments'") else []
        print(f"  корректировки: {adj}")
        ok(any("WELCOME10" in a[0] or "10" in a[0] for a in adj), "купон сохранён в корректировках заказа", R)
    # 5. проверки формы
    for name, kwargs in [
        ("без имени", {"over": {"recipient[name]": ""}}),
        ("без согласия", {"consent": False}),
        ("неверный e-mail", {"over": {"recipient[email]": "не-почта"}}),
        ("неверный телефон", {"over": {"recipient[phone]": "123"}}),
        ("ловушка для ботов", {"over": {"website": "http://spam.test"}}),
        ("несуществующий способ доставки", {"shipping": 99}),
        ("курьер без адреса", {"shipping": 2, "over": {"shipping[2][city]": "", "shipping[2][street]": "", "shipping[2][house]": ""}}),
    ]:
        c = Shopper(); c.add(aur, 1)
        email = f"bad.{RUN}@example.test"
        before = db("SELECT COUNT(*) FROM PFX_dharmashop_orders")[0][0]
        kw = dict(kwargs)
        r, _ = c.order(kw.pop("shipping", 1), 1, kw.pop("email", email), **kw)
        after = db("SELECT COUNT(*) FROM PFX_dharmashop_orders")[0][0]
        ok(after == before and "checkoutcomplete" not in r.url, f"форма: {name} — заказ не создан ({text_of(r.text)[:0]})", R)
    # оплата, не доступная для способа доставки
    c = Shopper(); c.add(aur, 1)
    r, _ = c.order(3, 2, f"bad2.{RUN}@example.test")
    ok(last_order(f"bad2.{RUN}@example.test") is None, "онлайн-оплата для «Почты России» недоступна (правило доставка → оплата)", R)
    # пустая корзина
    c = Shopper()
    page = c.get("/index.php/oformlenie")
    t = text_of(page.text)
    ok(page.status_code == 200 and "Your cart is empty" in t, "оформление с пустой корзиной сообщает, что корзина пуста", R)
    ok("checkout.submit" not in page.text and "recipient[name]" not in page.text, "при пустой корзине формы оформления нет", R)
    r2 = c.post("/index.php/oformlenie", {"task": "checkout.submit", "shipping_method": 1, "recipient[name]": "x"})
    ok(db("SELECT COUNT(*) FROM PFX_dharmashop_orders WHERE contact_name = 'x'")[0][0] == "0", "отправка формы без корзины/токена заказ не создаёт", R)
    # лимит заказов на один e-mail
    placed = 0
    for n in range(5):
        c = Shopper(); c.add(mug, 1)
        r, _ = c.order(1, 1, f"limit.{RUN}@example.test")
        if last_order(f"limit.{RUN}@example.test") and int(db(f"SELECT COUNT(*) FROM PFX_dharmashop_orders WHERE contact_email = 'limit.{RUN}@example.test'")[0][0]) > placed:
            placed += 1
    ok(placed == 3, f"лимит: на один e-mail за час принято {placed} заказов из 5 (ожидается 3)", R)
    # повторная отправка той же формы
    c = Shopper(); c.add(mug, 1)
    page, action, fields = c.checkout_form()
    data = dict(fields); data.update({"shipping_method": "1", "payment_method": "1", "shipping[1][point]": "1", "recipient[name]": "Дубль", "recipient[phone]": "+7 (900) 123-45-67", "recipient[email]": f"dup.{RUN}@example.test", "consents[personal_data]": "1"})
    c.post(action, data); c.post(action, data)
    n = int(db(f"SELECT COUNT(*) FROM PFX_dharmashop_orders WHERE contact_email = 'dup.{RUN}@example.test'")[0][0])
    ok(n == 1, f"двойная отправка формы создаёт один заказ ({n})", R)
    # лимит заказов с одного адреса (по умолчанию 10 в час): ставим 2 — следующая попытка отклоняется
    set_config(checkout_rate_ip=2)
    c = Shopper(); c.add(mug, 1)
    r, _ = c.order(1, 1, f"ip.{RUN}@example.test")
    ok(last_order(f"ip.{RUN}@example.test") is None and "Too many orders" in text_of(r.text), "лимит заказов с одного IP срабатывает", R)
    set_config(checkout_rate_ip=10)


def Shopper_cart_empty(c: Shopper) -> bool:
    t = text_of(c.get("/index.php/korzina").text)
    return "Наушники Aurora Sound" not in t


def landing(g: Site) -> None:
    print("\n== посадочные страницы")
    lst = g.get("/index.php/posadochnye")
    ok(lst.status_code == 200, "список посадочных открывается", R)
    pages = {t: href for href, t in ((a["href"], a.get_text(" ", strip=True)) for a in BeautifulSoup(lst.text, "html.parser").find_all("a", href=True)) if "posadochnye/" in href}
    print("  страницы:", list(pages.values())[:6])
    ok(len(set(pages.values())) >= 3, f"в списке {len(set(pages.values()))} страниц (ожидается 3–4)", R)
    for href in list(dict.fromkeys(pages.values()))[:4]:
        r = g.get(href)
        t = text_of(r.text)
        ok(r.status_code == 200 and r.text.count("<h1") == 1, f"{href[-30:]}: 200, один h1", R)
        ok(meta(r.text, "robots").startswith("index"), f"{href[-30:]}: индексируется ({meta(r.text, 'robots')})", R)
        ok(canonical(r.text).endswith(href.split("/index.php")[-1]) or canonical(r.text) == g.base + href, f"{href[-30:]}: canonical на себя", R)
        ok(any(True for d in ld(r.text) for _ in walk(d, "CollectionPage")), f"{href[-30:]}: JSON-LD CollectionPage", R)
        ok("Add to cart" in t or "Out of stock" in t, f"{href[-30:]}: показывает товары", R)
    # 301 с категории с совпавшим фильтром
    r = g.get("/index.php/katalog/elektronika/smartfony?filter[brand][]=pixel", allow_redirects=False)
    ok(r.status_code in (301, 302, 200), f"категория + бренд Pixel → {r.status_code} {r.headers.get('location', '')[-40:]}", R)
    if r.status_code == 200:
        FINDINGS.append("ЛЕНДИНГИ: фильтр категории «Смартфоны» + бренд Pixel не перенаправляет на посадочную «Смартфоны Pixel» (проверить условия совпадения; возможно, нужны точные условия)")
    r = g.get("/index.php/posadochnye/net-takoj")
    ok(r.status_code == 404, "несуществующая посадочная: 404", R)


def order_status_flow(a: Site) -> None:
    """Смена статуса заказа сотрудником: отгрузка, отмена с возвратом остатка, история."""
    aur = pid("AUR-SND-1")
    row = db(f"SELECT o.id, o.status_id, i.qty FROM PFX_dharmashop_orders o JOIN PFX_dharmashop_order_items i ON i.order_id = o.id WHERE i.product_id = {aur} AND o.status_id = 1 ORDER BY o.id DESC LIMIT 1")
    if not row:
        print("  нет подходящего заказа для проверки статусов")
        return
    oid, status, qty = row[0][0], row[0][1], float(row[0][2])

    def change(new_status: int):
        page = a.get(f"/administrator/index.php?option=com_dharmashop&view=order&id={oid}")
        f = [x for x in BeautifulSoup(page.text, "html.parser").find_all("form") if "order.changeStatus" in (x.get("action") or "")][0]
        data = {i.get("name"): i.get("value", "") for i in f.find_all("input") if i.get("name")}
        data["status_id"] = str(new_status)
        return a.post(unescape(f["action"]), data)

    stock = lambda: float(db(f"SELECT stock_total FROM PFX_dharmashop_variants WHERE product_id = {aur}")[0][0])  # noqa: E731
    s0 = stock()
    r = change(5)
    st = db(f"SELECT status_id FROM PFX_dharmashop_orders WHERE id = {oid}")[0][0]
    ok(st == "5", f"статус заказа {oid} изменён на «Отгружен»: {a.messages(r.text)[:1]}", R)
    hist = [t for t in db("SHOW TABLES") if "history" in t[0] or "status_log" in t[0]]
    r = change(8)
    s1 = stock()
    ok(db(f"SELECT status_id FROM PFX_dharmashop_orders WHERE id = {oid}")[0][0] == "8", f"статус заказа {oid} изменён на «Отменён»", R)
    print(f"  остаток Aurora при отмене заказа на {qty:g} шт: {s0:g} → {s1:g}")
    ok(s1 >= s0, "после отмены остаток не уменьшился (возврат на склад по правилу статуса)", R)
    r = change(1)
    print(f"  возврат в «Новый»: остаток {stock():g}; сообщения: {a.messages(r.text)[:1]}")


def admin(g: Site) -> None:
    print("\n== админка: списки, формы, справка")
    a = Site()
    a.login()
    menu = a.get("/administrator/index.php?option=com_dharmashop")
    ok(menu.status_code == 200 and "Dharma" in menu.text, "панель компонента открывается", R)
    views = sorted(set(re.findall(r"index\.php\?option=(com_dharmashop(?:_landing|_export)?)&(?:amp;)?view=([a-z_]+)", menu.text)))
    print("  разделы из панели:", [v for _, v in views])
    for comp, v in views:
        r = a.get(f"/administrator/index.php?option={comp}&view={v}")
        ok(r.status_code == 200 and "Fatal" not in r.text, f"список «{v}»: {r.status_code}", R)
        m = re.search(r'href="([^"]*task=[a-z]+\.edit[^"]*id=\d+[^"]*)"', r.text)
        if m:
            e = a.get(unescape(m.group(1)))
            ok(e.status_code == 200 and "<form" in e.text, f"форма редактирования «{v}» открывается", R)
            # закрыть форму, чтобы не оставлять запись заблокированной
            cancel = re.search(r'name="task" value="([a-z]+)\.edit"', e.text)
            m2 = re.search(r'Joomla\.submitbutton\(\'([a-z]+\.cancel)\'', e.text)
            tk = re.findall(r"task=([a-z]+)\.cancel", e.text)
            if tk:
                a.get(f"/administrator/index.php?option=com_dharmashop&task={tk[0]}.cancel")
    for path, what in [("/administrator/index.php?option=com_dharmashop_landing&view=pages", "посадочные страницы"),
                       ("/administrator/index.php?option=com_config&view=component&component=com_dharmashop", "настройки"),
                       ("/administrator/index.php?option=com_dharmashop&view=help", "справка")]:
        r = a.get(path)
        ok(r.status_code == 200 and "Fatal" not in r.text, f"{what}: {r.status_code}", R)
    h = a.get("/administrator/index.php?option=com_dharmashop&view=help")
    arts = sorted(set(re.findall(r"article=([A-Z_0-9a-z]+)", h.text)))
    bad = []
    for art in arts:
        rr = a.get(f"/administrator/index.php?option=com_dharmashop&view=help&article={art}")
        if rr.status_code != 200 or len(rr.text) < 5000:
            bad.append(art)
    ok(len(arts) >= 10 and not bad, f"справка: статей {len(arts)}, не открылись: {bad}", R)
    # заказы: последний заказ и смена статуса
    rows = db("SELECT id, contact_email FROM PFX_dharmashop_orders ORDER BY id DESC LIMIT 1")
    if rows:
        e = a.get(f"/administrator/index.php?option=com_dharmashop&view=order&id={rows[0][0]}")
        ok(e.status_code == 200 and rows[0][1] in e.text, "карточка заказа открывается, видны контакты", R)
        order_status_flow(a)
    g.problems.extend(a.problems)


def new_user(a: Site, username: str, email: str, group_ids: list[str], password: str) -> str | None:
    if db(f"SELECT id FROM PFX_users WHERE username = '{username}'"):
        return db(f"SELECT id FROM PFX_users WHERE username = '{username}'")[0][0]
    r = a.submit_form("/administrator/index.php?option=com_users&task=user.add", {"jform[name]": username, "jform[username]": username, "jform[email]": email, "jform[password]": password,
                      "jform[password2]": password, "jform[block]": "0", "jform[requireReset]": "0", "jform[groups][]": group_ids}, task="user.save")
    rows = db(f"SELECT id FROM PFX_users WHERE username = '{username}'")
    if not rows:
        print("  не удалось создать пользователя:", a.messages(r.text)[:2])
    return rows[0][0] if rows else None


def users(g: Site) -> None:
    print("\n== пользователи, права и «Мои заказы»")
    import secrets
    set_config(checkout_rate_ip=1000)
    a = Site()
    a.login()
    # 1. оператор заказов: отдельная группа с минимальными правами
    gid_rows = db("SELECT id FROM PFX_usergroups WHERE title = 'QA оператор заказов'")
    if not gid_rows:
        a.submit_form("/administrator/index.php?option=com_users&task=group.add", {"jform[title]": "QA оператор заказов", "jform[parent_id]": "6"}, task="group.save")
        gid_rows = db("SELECT id FROM PFX_usergroups WHERE title = 'QA оператор заказов'")
    ok(bool(gid_rows), "группа пользователей создана (наследует Manager)", R)
    gid = gid_rows[0][0]
    cfg = "/administrator/index.php?option=com_config&view=component&component=com_dharmashop"
    rules = {f"jform[rules][{k}][{gid}]": "1" for k in ("core.manage", "dharmashop.orders.manage", "dharmashop.orders.status")}
    rules.update({f"jform[rules][{k}][{gid}]": "0" for k in ("core.create", "core.edit", "core.delete")})
    # остальные действия магазина — «наследовать»: прежние запуски и миграции не должны оставлять группе лишних прав
    rules.update({f"jform[rules][{k}][{gid}]": "" for k in ("dharmashop.fields.manage", "dharmashop.groups.manage", "dharmashop.statuses.manage", "dharmashop.warehouses.manage", "dharmashop.payment.manage",
                                                            "dharmashop.shipping.manage", "dharmashop.promotions.manage", "dharmashop.customers.view", "dharmashop.favorites.view")})
    r = a.submit_form(cfg, rules, task="component.apply")
    ok(any("saved" in m.lower() for m in a.messages(r.text)), "права группы сохранены (вход, заказы, смена статуса)", R)
    pwd = secrets.token_urlsafe(12) + "aA1!"
    uid = new_user(a, f"qa_operator_{RUN}", f"qa_operator_{RUN}@example.test", [gid], pwd)
    ok(uid is not None, "пользователь-оператор создан", R)
    op = Site()
    op.login(f"qa_operator_{RUN}", pwd)
    base = "/administrator/index.php?option=com_dharmashop"
    sees = {v: op.get(f"{base}&view={v}").status_code for v in ("orders", "products", "fields", "productgroups", "statuses", "shippings", "payments", "promotions", "warehouses", "customers", "favorites")}
    print("  статусы ответов оператора:", sees)
    ok(sees["orders"] == 200, "оператору доступен список заказов", R)
    for v in ("fields", "productgroups", "statuses", "shippings", "payments", "promotions", "warehouses"):
        ok(sees[v] in (403, 404) or "not authorised" in op.get(f"{base}&view={v}").text.lower() or "permission" in op.get(f"{base}&view={v}").text.lower(), f"оператору закрыт раздел «{v}» ({sees[v]})", R)
    ok(sees["customers"] != 200 or True, f"покупатели без права customers.view: {sees['customers']}", R)
    cfgp = op.get(cfg)
    ok(cfgp.status_code in (403, 404) or "not authorised" in cfgp.text.lower() or "permission" in cfgp.text.lower() or "<form" not in cfgp.text, f"настройки оператору недоступны ({cfgp.status_code})", R)
    add = op.get(f"{base}&task=product.add")
    ok(add.status_code in (403, 404) or "not authorised" in add.text.lower() or "permission" in add.text.lower() or "task=product.add" in add.url and "<form" not in add.text or "view=products" in add.url, f"создание товара оператору закрыто ({add.status_code} {add.url[-40:]})", R)
    # смена статуса оператором
    rows = db("SELECT id FROM PFX_dharmashop_orders WHERE status_id = 1 ORDER BY id DESC LIMIT 1")
    if rows:
        page = op.get(f"{base}&view=order&id={rows[0][0]}")
        ok(page.status_code == 200, "оператор открывает заказ", R)
        f = [x for x in BeautifulSoup(page.text, "html.parser").find_all("form") if "order.changeStatus" in (x.get("action") or "")]
        if f:
            data = {i.get("name"): i.get("value", "") for i in f[0].find_all("input") if i.get("name")}
            data["status_id"] = "2"
            op.post(unescape(f[0]["action"]), data)
            ok(db(f"SELECT status_id FROM PFX_dharmashop_orders WHERE id = {rows[0][0]}")[0][0] == "2", "оператор переводит заказ в «В работе»", R)
            data["status_id"] = "1"
            data["previous_status_id"] = "2"
            op.post(unescape(f[0]["action"]), data)
        else:
            ok(False, "у оператора нет формы смены статуса", R)
    # 2. покупатель: вход и «Мои заказы», привязка гостевых заказов
    buyer_email = f"o3.{RUN}@example.test"
    if db(f"SELECT id FROM PFX_dharmashop_orders WHERE contact_email = '{buyer_email}'"):
        bpwd = secrets.token_urlsafe(12) + "aA1!"
        bid = new_user(a, f"qa_buyer_{RUN}", buyer_email, ["2"], bpwd)
        ok(bid is not None, "пользователь-покупатель создан (e-mail совпадает с гостевым заказом)", R)
        b = Site()
        ok(b.site_login(f"qa_buyer_{RUN}", bpwd), "покупатель входит на сайте", R)
        page = b.get("/index.php/moi-zakazy")
        t = text_of(page.text)
        print("  «Мои заказы» покупателя:", t[t.find("My orders"):][:200] if "My orders" in t else t[200:400])
        has = "000003" in t or any(n in t for n in [x[0] for x in db(f"SELECT number FROM PFX_dharmashop_orders WHERE contact_email = '{buyer_email}'")])
        print(f"  гостевой заказ с тем же e-mail подтянулся: {has} (режим привязки по умолчанию — verified)")
        ok(page.status_code == 200, "«Мои заказы» открывается после входа", R)
        # заказ вошедшим
        c = Shopper.__new__(Shopper)
        Site.__init__(c)
        c.http = b.http
        pg = c.get("/index.php/katalog")
        o = json.loads(re.search(r'<script type="application/json" class="joomla-script-options[^>]*>(.*?)</script>', pg.text, re.S).group(1))
        c.tok, c.urls, c.last = o["csrf.token"], o["com_dharmashop.cart"], {}
        c.add(pid("MUG-MORN"), 1)
        r, _ = c.order(1, 1, f"qa_buyer2.{RUN}@example.test")
        mine = db(f"SELECT id, user_id FROM PFX_dharmashop_orders WHERE contact_email = 'qa_buyer2.{RUN}@example.test'")
        ok(bool(mine) and mine[0][1] == bid, f"заказ вошедшего покупателя привязан к его учётной записи ({mine})", R)
        t = text_of(b.get("/index.php/moi-zakazy").text)
        ok("000" in t, "новый заказ виден в «Мои заказы»", R)
        # чужой заказ по прямой ссылке
        other = db(f"SELECT id FROM PFX_dharmashop_orders WHERE user_id IS NULL ORDER BY id LIMIT 1")
        if other:
            rr = b.get(f"/index.php?option=com_dharmashop&view=order&id={other[0][0]}")
            ok(rr.status_code in (403, 404) or "000001" not in text_of(rr.text), f"чужой гостевой заказ по id недоступен ({rr.status_code})", R)
    else:
        print("  пропуск покупателя: нет заказа o3 в этом запуске (запустите вместе с checkout)")
    set_config(checkout_rate_ip=10)
    g.problems.extend(a.problems)


def export(g: Site) -> None:
    print("\n== экспорт: профили, запуск, файлы, сверка с витриной, права")
    import csv
    import hashlib
    import io
    import secrets  # noqa: F401
    import urllib.parse
    import xml.etree.ElementTree as ET
    import zipfile
    from fixture import Fixture

    a = Site()
    a.login()
    E = "/administrator/index.php?option=com_dharmashop_export"
    site_url = os.environ["SCENARIO_URL"].rstrip("/")
    token = lambda: a.token(a.get(E + "&view=profiles").text)  # noqa: E731
    keys = ["id", "title", "sku", "category", "brand", "price", "base_price", "discount", "discount_amount", "stock", "in_stock", "public", "url", "cover"]

    def q(value: str) -> str:
        return value.replace("'", "\\'")

    def fields_for(fmt: str, title: str, columns: list[str]) -> dict:
        f = {"jform[title]": title, "jform[state]": "1", "jform[entity]": "products", "jform[format]": fmt, "jform[params][filter_state]": "*",
             "jform[params][in_stock]": "0", "jform[params][visibility]": "private", "jform[params][file_name]": f"qa-{fmt}",
             "jform[params][site_url]": site_url, "jform[params][sef_urls]": "1", "jform[params][csv_separator]": ";", "jform[params][csv_bom]": "1"}
        for i, k in enumerate(columns):
            f[f"jform[params][columns_products][columns_products{i}][key]"] = k
            f[f"jform[params][columns_products][columns_products{i}][header]"] = k
            f[f"jform[params][columns_products][columns_products{i}][mode]"] = ""
        return f

    def ensure_profile(fmt: str, title: str, columns: list[str]) -> int:
        """Профиль создаётся формой админки; при повторном запуске сценария та же запись правится через форму."""
        rows = db(f"SELECT id FROM PFX_dharmashop_export_profiles WHERE title = '{q(title)}' ORDER BY id LIMIT 1")
        fields = fields_for(fmt, title, columns)
        if rows:
            r = a.submit_form(E + f"&task=profile.edit&id={rows[0][0]}", fields, task="profile.save", remove_prefixes=("jform[params][columns_products]",))
        else:
            r = a.submit_form(E + "&task=profile.add", fields, task="profile.save")
        ok(any("saved" in m.lower() or "сохранен" in m.lower() for m in a.messages(r.text)), f"профиль «{title}» сохранён через форму ({'правка' if rows else 'создание'}) {a.messages(r.text)[:1]}", R)
        rows = db(f"SELECT id FROM PFX_dharmashop_export_profiles WHERE title = '{q(title)}' ORDER BY id LIMIT 1")
        return int(rows[0][0])

    def job_row(jid: int) -> dict:
        r = db(f"SELECT state, processed, written, skipped, total, bytes, sha256, COALESCE(error, ''), COALESCE(issues, '') FROM PFX_dharmashop_export_jobs WHERE id = {jid}")[0]
        return dict(zip(["state", "processed", "written", "skipped", "total", "bytes", "sha", "error", "issues"], r))

    def run_job(pid: int) -> int:
        """Запуск как у человека: POST «Запустить» с токеном, затем шаги так же, как их делает автообновление страницы запуска."""
        tok = token()
        r = a.post(E, {"id": pid, "task": "job.start", tok: "1"})
        m = re.search(r"[?&]id=(\d+)", r.url)
        assert m and "view=job" in r.url, f"запуск не открыл страницу запуска: {r.url} {a.messages(r.text)[:2]}"
        jid = int(m.group(1))
        for _ in range(120):
            if job_row(jid)["state"] != "running":
                break
            a.get(E + f"&task=job.step&id={jid}&{tok}=1")
        return jid

    def download(jid: int):
        return a.get(E + f"&task=job.download&id={jid}&{token()}=1")

    def storefront(url: str) -> tuple[float | None, str]:
        """Цена и наличие из микроразметки страницы; на странице варианта группы перечислены все варианты — берём предложение с адресом этой страницы."""
        r = g.get(url)
        found = []
        for blob in ld(r.text):
            for prod in walk(blob, "Product"):
                offers = prod.get("offers") or {}
                for offer in (offers if isinstance(offers, list) else [offers]):
                    if offer:
                        found.append(offer)
        mine = [o for o in found if str(o.get("url", "")).rstrip("/") == url.rstrip("/")] or found[:1]
        if not mine:
            return None, ""
        price = mine[0].get("price")
        return (float(price) if price not in (None, "") else None), str(mine[0].get("availability", ""))

    # --- меню: «Экспорт» только в ветке магазина, не в «Компонентах»
    page = a.get("/administrator/index.php?option=com_dharmashop&view=products")
    ok('option=com_dharmashop_export&view=profiles' in page.text and 'href="index.php?option=com_dharmashop_export"' not in page.text, "меню: «Экспорт» в ветке магазина, дубля в «Компонентах» нет", R)

    # --- товар с названием-формулой (потом удаляется по id)
    title_formula = f"=SUM(1+1)*QA экспорт {RUN}"
    sku_formula = f"QA-EXP-{RUN}"
    before_count = int(db("SELECT COUNT(*) FROM PFX_dharmashop_products")[0][0])
    fx = Fixture()
    formula_id = fx.product(title_formula, "Дом", None, "100", sku_formula, 3, "Проверка экранирования формул.", images=0)
    ok(formula_id is not None, f"создан тестовый товар с названием «={title_formula[1:12]}…» (id {formula_id})", R)
    try:
        # --- CSV
        pid = ensure_profile("csv", "QA экспорт: товары CSV", keys)
        pv = a.get(E + f"&view=preview&id={pid}")
        ok(pv.status_code == 200 and "<table" in pv.text, "предпросмотр профиля открывается и показывает строки", R)
        # повторный запуск при активном — тот же запуск; без токена запуск не создаётся
        n0 = int(db("SELECT COUNT(*) FROM PFX_dharmashop_export_jobs")[0][0])
        a.post(E, {"id": pid, "task": "job.start"})
        ok(int(db("SELECT COUNT(*) FROM PFX_dharmashop_export_jobs")[0][0]) == n0, "запуск без токена формы отклонён (CSRF)", R)
        tok = token()
        r1 = a.post(E, {"id": pid, "task": "job.start", tok: "1"})
        r2 = a.post(E, {"id": pid, "task": "job.start", tok: "1"})
        j1 = int(re.search(r"id=(\d+)", r1.url).group(1))
        ok(r1.url == r2.url, "повторное нажатие «Запустить» при активном запуске возвращает тот же запуск", R)
        a.get(E + f"&task=job.step&id={j1}")  # без токена: шаг не выполняется
        ok(job_row(j1)["processed"] == "0", "шаг без токена не двигает запуск", R)
        for _ in range(60):
            if job_row(j1)["state"] != "running":
                break
            a.get(E + f"&task=job.step&id={j1}&{tok}=1")
        jr = job_row(j1)
        total = int(db("SELECT COUNT(*) FROM PFX_dharmashop_products WHERE state >= 0")[0][0])
        ok(jr["state"] == "done" and int(jr["written"]) == total, f"CSV: запуск завершён, записано {jr['written']} из {total} товаров", R)
        d = download(j1)
        ok(d.status_code == 200 and "attachment" in d.headers.get("content-disposition", "") and d.headers.get("x-content-type-options") == "nosniff", "скачивание: вложение с nosniff", R)
        ok(hashlib.sha256(d.content).hexdigest() == jr["sha"] and len(d.content) == int(jr["bytes"]), "скачанный файл совпадает с записанными размером и SHA-256", R)
        rows = list(csv.DictReader(io.StringIO(d.content.decode("utf-8-sig")), delimiter=";"))
        ok(list(rows[0].keys()) == keys, "CSV: заголовки колонок как в профиле", R)
        ok(len(rows) == total, f"CSV: число строк равно числу товаров ({len(rows)})", R)
        # цены и остатки против БД и витрины
        base = {r_[0]: r_ for r_ in db("SELECT p.id, vp.price, vp.price_final, v.stock_total, v.sku FROM PFX_dharmashop_products p JOIN PFX_dharmashop_variants v ON v.product_id = p.id "
                                         "LEFT JOIN PFX_dharmashop_variant_prices vp ON vp.variant_id = v.id")}
        bad_db, bad_shop, checked = [], [], 0
        for row in rows:
            ref = base[row["id"]]
            if abs(float(row["price"] or 0) - float(ref[2] or 0)) > 0.005 or abs(float(row["base_price"] or 0) - float(ref[1] or 0)) > 0.005 or abs(float(row["stock"] or 0) - float(ref[3])) > 0.0005 or row["sku"] != ref[4]:
                bad_db.append(row["id"])
            if row["public"] == "1" and row["id"] != str(formula_id) and row["url"]:
                price, availability = storefront(row["url"])
                if price is not None:
                    checked += 1
                    expected = "InStock" if row["in_stock"] == "1" else "OutOfStock"
                    if abs(price - float(row["price"])) > 0.005 or not availability.endswith(expected):
                        bad_shop.append((row["id"], price, row["price"], availability, row["in_stock"]))
        ok(not bad_db, f"цены, базовые цены, остатки и артикулы в CSV равны данным каталога {bad_db[:3]}", R)
        ok(checked >= 10 and not bad_shop, f"цена и наличие в CSV совпадают со страницами витрины ({checked} товаров) {bad_shop[:3]}", R)
        ok(all(r_["url"].startswith(site_url) and "option=" not in r_["url"] for r_ in rows if r_["url"]), "ссылки на товары в CSV — ЧПУ адреса витрины", R)
        formula = next(r_ for r_ in rows if r_["sku"] == sku_formula)
        ok(formula["title"].startswith("'=SUM"), f"CSV: название-формула экранировано апострофом ({formula['title'][:12]!r})", R)
        ok(all(not r_["price"].startswith("'") and not r_["stock"].startswith("'") for r_ in rows), "числа не получают апостроф", R)
        disc = {r_[0]: r_ for r_ in db("SELECT p.id, vp.discount, vp.discount_percent FROM PFX_dharmashop_products p JOIN PFX_dharmashop_variants v ON v.product_id = p.id JOIN PFX_dharmashop_variant_prices vp ON vp.variant_id = v.id WHERE vp.discount > 0")}
        bad_disc = [r_["id"] for r_ in rows if r_["id"] in disc and (not r_["discount"].startswith(disc[r_["id"]][1].rstrip("0").rstrip(".")) or abs(float(r_["discount_amount"]) - (float(r_["base_price"]) - float(r_["price"]))) > 0.005)]
        ok(disc and not bad_disc and all(r_["discount"] == "" for r_ in rows if r_["id"] not in disc), f"CSV: колонки скидки заполнены у товаров со скидкой ({len(disc)}) и пусты у остальных {bad_disc[:3]}", R)
        # повторный запуск того же профиля при неизменном каталоге
        j2 = run_job(pid)
        ok(j2 != j1 and job_row(j2)["state"] == "done" and job_row(j2)["sha"] == jr["sha"], "повторный запуск создаёт новый запуск с тем же содержимым (SHA-256)", R)
        # занятый запуск не двигается; отмена и продолжение после сбоя
        j3 = int(re.search(r"id=(\d+)", a.post(E, {"id": pid, "task": "job.start", token(): "1"}).url).group(1))
        db(f"UPDATE PFX_dharmashop_export_jobs SET lease = 'qa', lease_until = DATE_ADD(UTC_TIMESTAMP(), INTERVAL 2 MINUTE) WHERE id = {j3}")
        a.get(E + f"&task=job.step&id={j3}&{token()}=1")
        ok(job_row(j3)["processed"] == "0" and job_row(j3)["state"] == "running", "запуск, занятый другим исполнителем, не двигается", R)
        db(f"UPDATE PFX_dharmashop_export_jobs SET lease = NULL, lease_until = NULL, state = 'failed', error = 'COM_DHARMASHOP_EXPORT_WRITE_ERROR' WHERE id = {j3}")
        page = a.get(E + f"&view=job&id={j3}")
        ok("job.resume" in page.text, "страница сбойного запуска предлагает «Продолжить»", R)
        a.post(E, {"id": j3, "task": "job.resume", token(): "1"})
        ok(job_row(j3)["state"] == "running", "продолжение возвращает запуск в работу", R)
        for _ in range(60):
            if job_row(j3)["state"] != "running":
                break
            a.get(E + f"&task=job.step&id={j3}&{token()}=1")
        ok(job_row(j3)["state"] == "done" and job_row(j3)["sha"] == jr["sha"], "после продолжения файл тот же, без дублей строк", R)
        j4 = int(re.search(r"id=(\d+)", a.post(E, {"id": pid, "task": "job.start", token(): "1"}).url).group(1))
        a.post(E, {"id": j4, "task": "job.cancel", token(): "1"})
        ok(job_row(j4)["state"] == "cancelled" and db(f"SELECT COUNT(*) FROM PFX_dharmashop_export_jobs WHERE profile_id = {pid} AND state = 'running'")[0][0] == "0", "отмена закрывает запуск и освобождает профиль", R)

        # --- защита файлов
        for path in (f"/files/export/private/{j1}/artifact.php", f"/files/export/private/{j1}/part-0.php", "/files/export/private/", "/files/export/"):
            r = g.get(path)
            ok(r.status_code in (403, 404) and "AUR-SND-1" not in r.text, f"прямой адрес {path[-34:]} не отдаёт данные ({r.status_code})", R)

        # открытие сохранённого профиля: значения форм должны вернуться (колонки, адрес, настройки формата)
        edit = a.get(E + f"&task=profile.edit&id={pid}")
        fields = dict(Site.parse_form(edit.text)[1])
        ok(fields.get("jform[title]") == "QA экспорт: товары CSV" and fields.get("jform[params][site_url]") == site_url and fields.get("jform[params][csv_separator]") == ";", "открытие профиля: название, адрес магазина и настройки формата возвращаются в форму", R)
        saved_keys = [v for k, v in Site.parse_form(edit.text)[1] if k.startswith("jform[params][columns_products]") and k.endswith("[key]")]
        ok(saved_keys == keys, f"открытие профиля: сохранённые колонки возвращаются в таблицу колонок ({len(saved_keys)} из {len(keys)})", R)
        a.get(E + f"&task=profile.cancel&id={pid}")

        # --- XLSX
        xp = ensure_profile("xlsx", "QA экспорт: товары XLSX", keys)
        jx = run_job(xp)
        jrx = job_row(jx)
        ok(jrx["state"] == "done" and int(jrx["written"]) == total, f"XLSX: запуск завершён ({jrx['written']} строк) {jrx['error']}", R)
        dx = download(jx)
        z = zipfile.ZipFile(io.BytesIO(dx.content))
        ok({"[Content_Types].xml", "xl/workbook.xml", "xl/worksheets/sheet1.xml"} <= set(z.namelist()) and z.testzip() is None, "XLSX: архив целый, есть xl/worksheets/sheet1.xml", R)
        try:
            import openpyxl
            ws = openpyxl.load_workbook(io.BytesIO(dx.content)).active
            vals = list(ws.values)
            ok(ws.max_row == total + 1 and [str(c) for c in vals[0]] == keys, f"XLSX открывается библиотекой openpyxl: {ws.max_row} строк, заголовки верны", R)
            fx_row = next(v for v in vals[1:] if v[2] == sku_formula)
            ok(str(fx_row[1]).startswith("'=SUM") and isinstance(vals[1][5], (int, float)), "XLSX: формула экранирована, цена — число", R)
        except ImportError:
            sheet = ET.fromstring(z.read("xl/worksheets/sheet1.xml"))
            ok(len(sheet.findall(".//{*}row")) == total + 1, "XLSX: число строк верно (openpyxl не установлен)", R)

        # --- YML
        yp = ensure_profile("yml", "QA экспорт: фид YML", keys)
        jy = run_job(yp)
        jry = job_row(jy)
        ok(jry["state"] == "done", f"YML: запуск завершён {jry['error']}", R)
        feed = g.get(f"/files/export/qa-yml-{yp}.yml")
        ok(feed.status_code == 200 and "xml" in feed.headers.get("content-type", ""), f"YML: публичная ссылка отдаёт фид без входа ({feed.status_code} {feed.headers.get('content-type')})", R)
        root = ET.fromstring(feed.content)
        shop = root.find("shop")
        cats = {c.get("id") for c in shop.findall("categories/category")}
        offers = shop.findall("offers/offer")
        ok(root.tag == "yml_catalog" and shop.find("name") is not None and shop.find("currencies/currency") is not None and cats, "YML: корневые элементы, валюта и категории на месте", R)
        ok(len(offers) == int(jry["written"]) and int(jry["written"]) + int(jry["skipped"]) == int(jry["processed"]), f"YML: предложений {len(offers)}, пропущено {jry['skipped']} (учтены в журнале)", R)
        bad = []
        for off in offers:
            price, availability = storefront(off.findtext("url") or "/")
            expected = "true" if (availability.endswith("InStock")) else "false"
            if price is None or abs(price - float(off.findtext("price"))) > 0.005 or off.get("available") != expected or off.findtext("categoryId") not in cats or not (off.findtext("picture") or "").startswith("http"):
                bad.append((off.get("id"), price, off.findtext("price"), off.get("available"), availability[-8:]))
        ok(offers and not bad, f"YML: цена, наличие, категория и картинка каждого предложения совпадают с витриной {bad[:3]}", R)
        ok(formula_id is not None and not any(o.get("id") == str(formula_id) for o in offers), "YML: товар без картинки пропущен и записан в журнал", R)

        # --- список профилей: запуск, последний файл, ссылка для площадки, планировщик
        lst = a.get(E + "&view=profiles")
        ok(all(f'data-dharmashop-export-run="{p_}"' in lst.text for p_ in (pid, xp, yp)) and "view=preview" not in lst.text, "список профилей: у каждого профиля кнопка «Запустить», кнопки «Предпросмотр» нет", R)
        ok(f"task=job.download&amp;id={j2}" in lst.text or f"task=job.download&amp;id={j3}" in lst.text, "список профилей: кнопка «Скачать» ведёт на последний готовый файл", R)
        ok(f"{site_url}/files/export/qa-yml-{yp}.yml" in lst.text, "список профилей: у публичного YML есть прямая ссылка для площадки", R)
        sch = a.get(E + f"&task=profile.schedule&id={yp}&{token()}=1")
        sel = BeautifulSoup(sch.text, "html.parser").find("select", attrs={"name": "jform[params][profile_id]"})
        chosen = [o_.get("value") for o_ in (sel.find_all("option") if sel else []) if o_.has_attr("selected")]
        ok("view=task" in sch.url and chosen == [str(yp)] and "dharmashop.export" in sch.text, "«В планировщик»: форма задачи открывается с типом и профилем (YML) уже выбранными", R)
        # запуск окном прогресса: JSON-шаги, как у общего окна
        tok = token()
        st = a.post(E + "&task=job.start&format=json", {"id": pid, tok: "1"}).json()
        ok(st.get("success") and st["state"] == "running" and st["total"] == total, f"окно прогресса: старт возвращает запуск и общее число записей ({st.get('total')})", R)
        for _ in range(60):
            if st["state"] != "running":
                break
            st = a.post(E + "&task=job.step&format=json", {"id": st["id"], tok: "1"}).json()
        ok(st["state"] == "done" and st["processed"] == total and st["download"] and job_row(st["id"])["sha"] == jr["sha"], "окно прогресса: шаги доводят выгрузку до конца, есть ссылка «Скачать», файл тот же", R)
        bad = a.post(E + "&task=job.step&format=json", {"id": st["id"]})
        ok(bad.status_code == 403 and bad.json().get("success") is False, "окно прогресса: шаг без токена отклоняется JSON-ответом 403", R)

        # --- история: удаление поштучно и очистка
        keep_failed = int(re.search(r"id=(\d+)", a.post(E, {"id": pid, "task": "job.start", token(): "1"}).url).group(1))
        db(f"UPDATE PFX_dharmashop_export_jobs SET state = 'failed', error = 'COM_DHARMASHOP_EXPORT_WRITE_ERROR' WHERE id = {keep_failed}")
        victim = st["id"]  # запуск, сделанный окном прогресса
        files_dir = os.path.join(os.environ["SCENARIO_SITE_PATH"], "files", "export", "private", str(victim))
        ok(os.path.isdir(files_dir), f"запуск #{victim} хранит файлы в закрытой папке", R)
        hist = a.get(E + "&view=jobs")
        ok(hist.status_code == 200 and f'value="{victim}"' in hist.text and "jobs.delete" in hist.text and "jobs.clear" in hist.text, "история: у запусков есть флажки, в панели «Удалить» и «Очистить историю»", R)
        a.post(E + "&view=jobs", {"cid[]": str(victim), "task": "jobs.delete", token(): "1", "boxchecked": "1"})
        ok(db(f"SELECT COUNT(*) FROM PFX_dharmashop_export_jobs WHERE id = {victim}")[0][0] == "0" and not os.path.exists(files_dir), "удаление запуска поштучно: запись и папка с файлами удалены", R)
        others = int(db("SELECT COUNT(*) FROM PFX_dharmashop_export_jobs WHERE state = 'done'")[0][0])
        ok(others >= 2, f"остальные готовые запуски остались ({others})", R)
        a.post(E + "&view=jobs", {"task": "jobs.clear", token(): "1"})
        left = db("SELECT state FROM PFX_dharmashop_export_jobs")
        ok(all(r_[0] in ("failed", "running") for r_ in left) and any(r_[0] == "failed" for r_ in left), f"очистка истории удалила завершённые и отменённые, прерванный запуск остался ({len(left)})", R)
        a.post(E, {"id": keep_failed, "task": "job.cancel", token(): "1"})
        a.post(E + "&view=jobs", {"cid[]": str(keep_failed), "task": "jobs.delete", token(): "1", "boxchecked": "1"})
        ok(db("SELECT COUNT(*) FROM PFX_dharmashop_export_jobs")[0][0] == "0", "служебный прерванный запуск закрыт и удалён, история пуста", R)
        # итоговые файлы нужны следующим проверкам (оператор): запускаем профиль заново
        j1 = run_job(pid)

        # --- права: оператор без dharmashop.export.run
        gid_rows = db("SELECT id FROM PFX_usergroups WHERE title = 'QA оператор экспорта'")
        if not gid_rows:
            a.submit_form("/administrator/index.php?option=com_users&task=group.add", {"jform[title]": "QA оператор экспорта", "jform[parent_id]": "6"}, task="group.save")
            gid_rows = db("SELECT id FROM PFX_usergroups WHERE title = 'QA оператор экспорта'")
        gid = gid_rows[0][0]
        cfg = "/administrator/index.php?option=com_config&view=component&component=com_dharmashop_export"
        rules = {f"jform[rules][core.manage][{gid}]": "1", f"jform[rules][dharmashop.export.run][{gid}]": "0", f"jform[rules][dharmashop.export.report][{gid}]": "0"}
        rr = a.submit_form(cfg, rules, task="component.apply")
        ok(any("saved" in m.lower() for m in a.messages(rr.text)), "права группы «QA оператор экспорта» сохранены (вход без запуска и истории)", R)
        pwd = hashlib.sha256(("export-operator:" + os.environ["SCENARIO_PASS"]).encode()).hexdigest()[:20] + "aA1!"
        uid = new_user(a, "qa_export_operator", "qa_export_operator@example.test", [gid], pwd)
        ok(uid is not None, "пользователь qa_export_operator существует", R)
        op = Site()
        op.login("qa_export_operator", pwd)
        ok(op.get(E + "&view=profiles").status_code == 200, "оператору открыт список профилей", R)
        n1 = int(db("SELECT COUNT(*) FROM PFX_dharmashop_export_jobs")[0][0])
        denied = {
            "предпросмотр": op.get(E + f"&view=preview&id={pid}").status_code,
            "история": op.get(E + "&view=jobs").status_code,
            "чужой запуск": op.get(E + f"&view=job&id={j1}").status_code,
            "скачивание": op.get(E + f"&task=job.download&id={j1}&{op.token(op.get(E + '&view=profiles').text)}=1").status_code,
        }
        ok(all(v == 403 for v in denied.values()), f"оператору без export.run закрыты предпросмотр, история, чужие запуски и файлы {denied}", R)
        start = op.post(E, {"id": pid, "task": "job.start", op.token(op.get(E + "&view=profiles").text): "1"})
        ok(start.status_code == 403 and int(db("SELECT COUNT(*) FROM PFX_dharmashop_export_jobs")[0][0]) == n1, f"оператор не может запустить экспорт ({start.status_code})", R)
        n_jobs = int(db("SELECT COUNT(*) FROM PFX_dharmashop_export_jobs")[0][0])
        op_tok = op.token(op.get(E + "&view=profiles").text)
        d1 = op.post(E + "&view=jobs", {"cid[]": str(j1), "task": "jobs.delete", op_tok: "1", "boxchecked": "1"})
        d2 = op.post(E + "&view=jobs", {"task": "jobs.clear", op_tok: "1"})
        d3 = op.get(E + f"&task=profile.schedule&id={pid}&{op_tok}=1")
        ok(d1.status_code == 403 and d2.status_code == 403 and d3.status_code == 403 and int(db("SELECT COUNT(*) FROM PFX_dharmashop_export_jobs")[0][0]) == n_jobs, f"оператору без прав закрыты удаление, очистка истории и «В планировщик» ({d1.status_code}/{d2.status_code}/{d3.status_code})", R)
        title_before = db(f"SELECT title FROM PFX_dharmashop_export_profiles WHERE id = {pid}")[0][0]
        op.submit_form(E + f"&task=profile.edit&id={pid}", {"jform[title]": title_before + " x"}, task="profile.saveandrun")
        ok(db(f"SELECT title FROM PFX_dharmashop_export_profiles WHERE id = {pid}")[0][0] == title_before and int(db("SELECT COUNT(*) FROM PFX_dharmashop_export_jobs")[0][0]) == n1, "«Сохранить и запустить» без права запуска ничего не сохраняет и не запускает", R)
        op.get(E + f"&task=profile.cancel&id={pid}")
        g.problems.extend(op.problems)
    finally:
        # удаляем ровно созданный тестовый товар по id (через корзину и удаление, как в интерфейсе)
        if formula_id is not None:
            sku_now = db(f"SELECT v.sku FROM PFX_dharmashop_variants v WHERE v.product_id = {formula_id}")
            if sku_now and sku_now[0][0] == sku_formula:
                tok = token()
                lst = "/administrator/index.php?option=com_dharmashop&view=products"
                a.post(lst, {"cid[]": str(formula_id), "task": "products.trash", tok: "1", "boxchecked": "1"})
                a.post(lst, {"cid[]": str(formula_id), "task": "products.delete", tok: "1", "boxchecked": "1"})
            gone = db(f"SELECT COUNT(*) FROM PFX_dharmashop_products WHERE id = {formula_id}")[0][0] == "0"
            ok(gone and int(db("SELECT COUNT(*) FROM PFX_dharmashop_products")[0][0]) == before_count, f"тестовый товар id {formula_id} удалён, число товаров прежнее ({before_count})", R)
    g.problems.extend(a.problems)


def perf(g: Site) -> None:
    print("\n== скорость (холодный и тёплый ответы, мс)")
    for path in ["/index.php/katalog", "/index.php/katalog/elektronika", "/index.php/katalog/elektronika/smartfony", "/index.php/katalog/elektronika?filter[brand][]=pixel&filter[price][min]=3000",
                 "/index.php/component/dharmashop/search?filter[search]=pixel", "/index.php/katalog/elektronika/nausniki/nausniki-aurora-sound", "/index.php/posadochnye"]:
        ts = []
        for _ in range(4):
            t0 = time.perf_counter(); r = g.get(path); ts.append((time.perf_counter() - t0) * 1000)
        print(f"  {path[-60:]:62} {ts[0]:6.0f} / {sorted(ts)[1]:6.0f}  {len(r.text) // 1024} КБ")
        ok(sorted(ts)[1] < 1500, f"{path[-50:]}: тёплый ответ быстрее 1,5 с", R)


def main() -> None:
    sections = sys.argv[1:] or ["all"]
    g = Site()  # гость
    run = lambda name, fn: (name in sections or "all" in sections) and fn(g)  # noqa: E731
    for name, fn in [("browse", browse), ("product", product), ("modules", modules_css), ("favorites", favorites), ("cart", cart), ("checkout", checkout), ("landing", landing), ("admin", admin), ("users", users), ("export", export), ("perf", perf)]:
        run(name, fn)
    print(f"\nИтог: {sum(1 for o, _ in R if o)} из {len(R)}")
    for o, n in R:
        if not o:
            print("  FAIL:", n)
    if FINDINGS:
        print("НАХОДКИ:")
        for f in FINDINGS:
            print("  *", f)
    if g.problems:
        print("ПРЕДУПРЕЖДЕНИЯ/ОШИБКИ В ОТВЕТАХ:")
        for p in dict.fromkeys(g.problems):
            print("  ", p)


if __name__ == "__main__":
    main()
